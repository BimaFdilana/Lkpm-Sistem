"""Normalisasi Excel DP.Proyek dan LKPM Non-UMK menjadi JSON Lines.

Laravel bertanggung jawab untuk menyimpan hasil ke MySQL. Worker ini hanya
membaca Excel secara streaming sehingga file besar tidak harus dimuat penuh ke
memori PHP.
"""

from __future__ import annotations

import argparse
import csv
import json
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook


PROJECT_HEADERS = {"Id Proyek", "Nib", "Nama Perusahaan", "Jumlah Investasi"}
LKPM_HEADERS = {"NO KODE PROYEK", "PERIODE LAPORAN", "TAHUN LAPORAN", "STATUS LAPORAN"}
SECTOR_HEADERS = {"Sektor Ekonomi", "Jumlah_KBLI"}


def text(value: Any) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def identifier(value: Any) -> str | None:
    value = text(value)
    if value is None:
        return None
    return value[:-2] if value.endswith(".0") else value


def project_identifier(value: Any) -> str | None:
    """Return a canonical project code for matching DP.Proyek with LKPM.

    LKPM exports numeric project codes with hyphens (for example
    ``201912-2915-0703-4686-697``), while DP.Proyek stores the same ID with an
    ``R-`` prefix (``R-201912291507034686697``). The source value is retained
    in ``raw``; only the matching key is compacted.
    """
    value = identifier(value)
    if value is None:
        return None

    value = value.replace("-", "")

    # DP.Proyek prefixes old OSS project IDs with "R-" while LKPM does not.
    # Example: R-201912291507034686697 == 201912-2915-0703-4686-697.
    if value.startswith("R") and value[1:].isdigit():
        return value[1:]

    return value


def number(value: Any) -> int:
    if value is None or value == "":
        return 0
    try:
        return int(Decimal(str(value).replace(",", "")))
    except (InvalidOperation, ValueError):
        return 0


def serializable(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def find_header_row(worksheet: Any, required: set[str]) -> tuple[int, list[str]]:
    for row_number, row in enumerate(worksheet.iter_rows(max_row=10, values_only=True), start=1):
        headers = [text(value) or "" for value in row]
        if required.issubset(set(headers)):
            return row_number, headers
    raise ValueError(f"Header wajib tidak ditemukan: {', '.join(sorted(required))}")


def rows_from_workbook(path: Path, required: set[str]) -> Iterator[tuple[str, dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    for worksheet in workbook.worksheets:
        header_row, headers = find_header_row(worksheet, required)
        for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            record = {
                header: serializable(row[index]) if index < len(row) else None
                for index, header in enumerate(headers)
                if header
            }
            if any(value is not None and value != "" for value in record.values()):
                yield worksheet.title, record


def rows_from_csv(path: Path) -> Iterator[tuple[str, dict[str, Any]]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for record in csv.DictReader(handle):
            yield path.stem, record


def project_record(record: dict[str, Any]) -> dict[str, Any] | None:
    project_code = project_identifier(record.get("Id Proyek"))
    nib = identifier(record.get("Nib"))
    if project_code is None or nib is None:
        return None

    return {
        "nib": nib,
        "company_name": text(record.get("Nama Perusahaan")) or "Tanpa nama perusahaan",
        "investment_status": text(record.get("Uraian Status Penanaman Modal")),
        "business_scale": text(record.get("Uraian Skala Usaha")),
        "address": text(record.get("Alamat Usaha")),
        "district": text(record.get("kecamatan_usaha")),
        "subdistrict": text(record.get("kelurahan_usaha")),
        "project_code": project_code,
        "project_name": text(record.get("nama_proyek")),
        "kbli": identifier(record.get("Kbli")),
        "kbli_description": text(record.get("Judul Kbli")),
        "sector": text(record.get("KL/Sektor Pembina")),
        "project_stage": text(record.get("Uraian_Jenis_Proyek")),
        "issued_at": text(record.get("Tanggal Terbit Oss")),
        "planned_investment": number(record.get("Jumlah Investasi")),
        "planned_tki": number(record.get("TKI")),
        "raw": record,
    }


def lkpm_record(record: dict[str, Any]) -> dict[str, Any] | None:
    project_code = project_identifier(record.get("NO KODE PROYEK"))
    year = number(record.get("TAHUN LAPORAN"))
    quarter = text(record.get("PERIODE LAPORAN"))
    if project_code is None or year == 0 or quarter is None:
        return None

    return {
        "project_code": project_code,
        "report_number": text(record.get("NO LAPORAN")),
        "report_year": year,
        "report_quarter": quarter,
        "reported_at": text(record.get("TANGGAL LAPORAN")),
        "report_status": text(record.get("STATUS LAPORAN")) or "Tidak diketahui",
        "total_investment_plan": number(record.get("NILAI TOTAL INVESTASI RENCANA")),
        "additional_investment": number(record.get("TOTAL TAMBAHAN INVESTASI")),
        "accumulated_investment": number(record.get("AKUMULASI REALISASI INVESTASI")),
        "accumulated_fixed_capital": number(record.get("AKUMULASI REALISASI MODAL TETAP")),
        "capital_explanation": text(record.get("PENJELASAN MODAL TETAP")),
        "planned_tki": number(record.get("JUMLAH RENCANA TKI")),
        "realized_tki": number(record.get("JUMLAH REALISASI TKI")),
        "planned_tka": number(record.get("JUMLAH RENCANA TKA")),
        "realized_tka": number(record.get("JUMLAH REALISASI TKA")),
        "raw": record,
    }


def sector_record(record: dict[str, Any]) -> dict[str, Any] | None:
    sector = text(record.get("Sektor Ekonomi"))
    if sector is None:
        return None
    return {"sector": sector, "kbli_count": number(record.get("Jumlah_KBLI")), "raw": record}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source_type", choices=("projects", "lkpm", "sectors"))
    parser.add_argument("file", type=Path)
    parser.add_argument("output", type=Path)
    arguments = parser.parse_args()

    required, normalizer = {
        "projects": (PROJECT_HEADERS, project_record),
        "lkpm": (LKPM_HEADERS, lkpm_record),
        "sectors": (SECTOR_HEADERS, sector_record),
    }[arguments.source_type]
    rows = rows_from_csv(arguments.file) if arguments.file.suffix.lower() == ".csv" else rows_from_workbook(arguments.file, required)

    arguments.output.parent.mkdir(parents=True, exist_ok=True)
    accepted = 0
    rejected = 0
    with arguments.output.open("w", encoding="utf-8") as handle:
        for sheet, raw_record in rows:
            normalized = normalizer(raw_record)
            if normalized is None:
                rejected += 1
                continue
            normalized["source_sheet"] = sheet
            handle.write(json.dumps(normalized, ensure_ascii=False, default=serializable) + "\n")
            accepted += 1

    print(json.dumps({"accepted_rows": accepted, "rejected_rows": rejected, "output": str(arguments.output)}))


if __name__ == "__main__":
    main()
